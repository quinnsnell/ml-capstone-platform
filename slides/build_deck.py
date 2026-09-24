#!/usr/bin/env python3
"""Build Lecture 2: CI/CD, Cloud Deployment, and Infrastructure as Code.

Starts from Lecture1_MLOps_Intro.pptx so the theme, fonts, and layouts match
exactly, then removes its slides and builds new ones.
"""
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = os.path.expanduser("~/Downloads/Lecture1_MLOps_Intro.pptx")
OUT = os.path.expanduser("~/classes/BigData/ml-capstone-platform/slides/Lecture2_CICD_IaC.pptx")

# theme palette from the source deck
DARK   = RGBColor(0x1F, 0x49, 0x7D)
BLUE   = RGBColor(0x4F, 0x81, 0xBD)
RED    = RGBColor(0xC0, 0x50, 0x4D)
GREEN  = RGBColor(0x9B, 0xBB, 0x59)
PURPLE = RGBColor(0x80, 0x64, 0xA2)
CYAN   = RGBColor(0x4B, 0xAC, 0xC6)
ORANGE = RGBColor(0xF7, 0x96, 0x46)
GREY   = RGBColor(0xD8, 0xD8, 0xD8)
LGREY  = RGBColor(0xF2, 0xF2, 0xF2)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
BLACK  = RGBColor(0x26, 0x26, 0x26)

prs = Presentation(SRC)

# --- strip the source slides, keep masters/layouts -------------------------
xml_slides = prs.slides._sldIdLst
for sld in list(xml_slides):
    rId = sld.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    prs.part.drop_rel(rId)
    xml_slides.remove(sld)

L_TITLE   = prs.slide_masters[0].slide_layouts[0]
L_CONTENT = prs.slide_masters[0].slide_layouts[1]
L_SECTION = prs.slide_masters[0].slide_layouts[2]
L_BLANK   = prs.slide_masters[0].slide_layouts[6]
L_TONLY   = prs.slide_masters[0].slide_layouts[5]


def title_slide(title, subtitle):
    s = prs.slides.add_slide(L_TITLE)
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle
    return s


def section(title, subtitle=""):
    s = prs.slides.add_slide(L_SECTION)
    s.shapes.title.text = title
    try:
        s.placeholders[1].text = subtitle
    except (KeyError, IndexError):
        pass
    return s


def bullets(title, items, size=20):
    """items: list of str, or (str, level) tuples."""
    s = prs.slides.add_slide(L_CONTENT)
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.clear()
    first = True
    for it in items:
        text, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = text
        p.level = lvl
        for r in p.runs:
            r.font.size = Pt(size if lvl == 0 else size - 3)
            if lvl > 0:
                r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return s


def code_slide(title, lines, size=13, caption=None):
    s = prs.slides.add_slide(L_TONLY)
    s.shapes.title.text = title
    top = Inches(1.6)
    h = Inches(5.0 if not caption else 4.4)
    box = s.shapes.add_textbox(Inches(0.5), top, Inches(9.0), h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = ln
        for r in p.runs:
            r.font.name = "Consolas"
            r.font.size = Pt(size)
            r.font.color.rgb = DARK if ln and not ln.startswith((" ", "#")) else BLACK
            if ln.strip().startswith("#"):
                r.font.color.rgb = RGBColor(0x70, 0x80, 0x90)
    if caption:
        cap = s.shapes.add_textbox(Inches(0.5), Inches(6.15), Inches(9.0), Inches(0.7))
        ctf = cap.text_frame; ctf.word_wrap = True
        ctf.text = caption
        for r in ctf.paragraphs[0].runs:
            r.font.size = Pt(15); r.font.color.rgb = DARK; r.font.italic = True
    return s


# ---- diagram primitives ---------------------------------------------------
def box(s, l, t, w, h, text="", fill=None, line=DARK, size=11, bold=False,
        color=BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, align=PP_ALIGN.CENTER):
    sh = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return sh


def label(s, l, t, w, h, text, size=10, color=DARK, bold=False, align=PP_ALIGN.LEFT, italic=False):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.italic = italic
    return tb


def arrow(s, l, t, w, h, color=BLUE, shape=MSO_SHAPE.RIGHT_ARROW):
    sh = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh

# ===========================================================================
# PART 0 — opening
# ===========================================================================
title_slide("CI/CD, the Cloud, and Infrastructure as Code",
            "ML Capstone · Lecture 2\nFrom a container on your laptop to a URL users can visit")

bullets("Where You Are Now", [
    "You can build an image and run a container",
    "You can push code to GitHub",
    "Your app runs — on your machine",
    "Missing: how does it get to a user?",
])

bullets("The Gap", [
    '"It works on my machine" is where we started',
    '"It works for users" is a different problem',
    "Someone has to build it — every time, the same way",
    "Someone has to run it — somewhere always on",
    "Someone has to decide it is safe to replace what is live",
    "Doing that by hand does not survive contact with a semester",
])

bullets("What We Are Building", [
    "One repo per student, two live environments",
    "Push code → tests run → app deploys → users see it",
    "Nobody clicks anything to deploy",
    "The same pipeline professionals use, scaled down",
])

# ===========================================================================
# ARCHITECTURE DIAGRAM
# ===========================================================================
s = prs.slides.add_slide(L_TONLY)
s.shapes.title.text = "The System, End to End"

# --- panel: your laptop
box(s, 0.25, 1.70, 2.70, 3.30, "", fill=LGREY, line=DARK, shape=MSO_SHAPE.RECTANGLE)
label(s, 0.32, 1.74, 2.56, 0.26, "YOUR LAPTOP", size=10, bold=True, align=PP_ALIGN.CENTER)
box(s, 0.40, 2.06, 2.40, 0.40, "editor  +  git", fill=WHITE, size=10)
box(s, 0.40, 2.58, 2.40, 1.52, "", fill=WHITE, line=BLUE)
label(s, 0.46, 2.62, 2.28, 0.24, "Docker Desktop", size=9, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
for i, (nm, wid) in enumerate([("hello", 0.80), ("time", 0.72), ("db", 0.60)]):
    xs = [0.52, 1.40, 2.18][i]
    box(s, xs, 2.94, wid, 0.36, nm, fill=BLUE, line=None, size=9, color=WHITE)
label(s, 0.46, 3.40, 2.28, 0.60,
      "docker compose up\nsame 3 services as production", size=8.5,
      color=BLACK, align=PP_ALIGN.CENTER)
label(s, 0.32, 4.20, 2.56, 0.70, "You test here first.\nFast, free, nobody sees it.",
      size=9, italic=True, align=PP_ALIGN.CENTER)

# --- panel: github
box(s, 3.55, 1.70, 2.55, 3.30, "", fill=LGREY, line=DARK, shape=MSO_SHAPE.RECTANGLE)
label(s, 3.60, 1.74, 2.45, 0.26, "GITHUB", size=10, bold=True, align=PP_ALIGN.CENTER)
box(s, 3.68, 2.06, 2.30, 0.38, "repo:  yourname-hello", fill=WHITE, size=9)
box(s, 3.68, 2.52, 1.08, 0.34, "staging", fill=ORANGE, line=None, size=9, color=WHITE)
box(s, 4.90, 2.52, 1.08, 0.34, "main", fill=GREEN, line=None, size=9, color=WHITE)
box(s, 3.68, 2.98, 2.30, 1.56, "", fill=WHITE, line=PURPLE)
label(s, 3.74, 3.02, 2.18, 0.24, "GitHub Actions  (CI)", size=9, bold=True,
      color=PURPLE, align=PP_ALIGN.CENTER)
box(s, 3.78, 3.32, 2.10, 0.32, "run tests", fill=PURPLE, line=None, size=9, color=WHITE)
box(s, 3.78, 3.70, 2.10, 0.32, "deploy-staging", fill=ORANGE, line=None, size=8.5, color=WHITE)
box(s, 3.78, 4.08, 2.10, 0.32, "deploy-prod", fill=GREEN, line=None, size=8.5, color=WHITE)
label(s, 3.60, 4.58, 2.45, 0.36, "tests gate both deploys", size=9, italic=True,
      align=PP_ALIGN.CENTER)

# --- panel: cloud
box(s, 6.70, 1.70, 3.05, 4.55, "", fill=LGREY, line=DARK, shape=MSO_SHAPE.RECTANGLE)
label(s, 6.75, 1.74, 2.95, 0.26, "CLOUD HOST — rigel.cs.byu.edu", size=9.5, bold=True,
      align=PP_ALIGN.CENTER)
label(s, 6.78, 2.02, 2.89, 0.24, "Coolify — builds images, runs containers", size=8.5,
      align=PP_ALIGN.CENTER)

for (ttl, ty, col) in [("STAGING", 2.32, ORANGE), ("PRODUCTION", 3.58, GREEN)]:
    box(s, 6.85, ty, 2.75, 1.14, "", fill=WHITE, line=col)
    label(s, 6.90, ty + 0.04, 2.64, 0.24, ttl, size=9, bold=True, color=col, align=PP_ALIGN.CENTER)
    for i, (nm, wid) in enumerate([("hello", 0.80), ("time", 0.72), ("db", 0.60)]):
        xs = [6.98, 7.86, 8.64][i]
        box(s, xs, ty + 0.34, wid, 0.34, nm, fill=col, line=None, size=9, color=WHITE)
    label(s, 6.90, ty + 0.74, 2.64, 0.28, "own database volume — never shared", size=8,
          italic=True, align=PP_ALIGN.CENTER)

box(s, 6.85, 4.86, 2.75, 0.44, "Traefik  —  routes by hostname", fill=CYAN, line=None,
    size=9, color=WHITE)
label(s, 6.78, 5.38, 2.89, 0.70,
      "yourname-hello-staging.ml-capstone.cs.byu.edu\nyourname-hello.ml-capstone.cs.byu.edu",
      size=8, align=PP_ALIGN.CENTER)

# --- users
box(s, 7.35, 6.62, 1.75, 0.48, "USERS", fill=DARK, line=None, size=10, color=WHITE)
label(s, 6.70, 7.08, 3.05, 0.30, "browser, on the CS VPN", size=8.5, align=PP_ALIGN.CENTER)
arrow(s, 8.06, 6.28, 0.32, 0.30, color=DARK, shape=MSO_SHAPE.UP_ARROW)

# --- arrows between panels
arrow(s, 2.99, 2.52, 0.52, 0.34)
label(s, 2.83, 2.22, 0.84, 0.26, "git push", size=8.5, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 6.14, 2.52, 0.52, 0.34, color=RED)
label(s, 5.94, 2.14, 0.92, 0.40, "deploy\nwebhook", size=8, bold=True, color=RED,
      align=PP_ALIGN.CENTER)

bullets("Your Laptop", [
    "Docker Compose runs the same three services production runs",
    ("hello — your app, the only one users reach", 1),
    ("time — a sidecar; stands in for a worker or model server", 1),
    ("db — Postgres, your data", 1),
    "If it does not work here, it will not work there",
    "Fast feedback: seconds, not a deploy cycle",
])

bullets("GitHub", [
    "Holds the code — and the history of every change",
    "Two branches, two meanings:",
    ("staging — what you are trying", 1),
    ("main — what users get", 1),
    "GitHub Actions is a computer that runs your tests on every push",
    "It is the only thing allowed to trigger a deploy",
])

bullets("The Cloud Host", [
    "Coolify is a PaaS — a layer over Docker that manages deploys",
    "It clones your repo, builds the image, starts the containers",
    "Two environments, fully separate, including databases",
    "Traefik sits in front and routes by hostname",
    ("one machine, many students, no port collisions", 1),
])

bullets("Users", [
    "They type a hostname — they never see containers",
    "Wildcard DNS: *.ml-capstone.cs.byu.edu → the class host",
    "Your repo name becomes your hostname",
    "On the CS VPN, because this is a classroom cluster",
])

bullets("Why Two Environments", [
    "Production is what users see. It should be boring.",
    "Staging is where you find out you were wrong",
    "Same image, same compose file, different data and URL",
    "Nothing reaches production without passing through staging",
    "This is not a classroom invention — it is standard practice",
])

# ===========================================================================
# PART B — CI/CD
# ===========================================================================
section("Continuous Integration & Continuous Deployment",
        "What happens between 'git push' and a user seeing it")

bullets("CI and CD", [
    "Continuous Integration — every change is built and tested, automatically",
    ("catches the break in minutes, not next Tuesday", 1),
    "Continuous Deployment — every change that passes goes live, automatically",
    ("the deploy is boring because it happens constantly", 1),
    "The point is not speed. The point is that it is the same every time.",
])

# --- CI/CD flow diagram ----------------------------------------------------
s = prs.slides.add_slide(L_TONLY)
s.shapes.title.text = "One Push, Step by Step"

STEPS = [("1\npush to\nstaging", BLUE), ("2\nActions\nruns tests", PURPLE),
         ("3\nPOST deploy\nwebhook", RED), ("4\nCoolify builds\n+ starts", CYAN),
         ("5\n/health\nchecks out", ORANGE)]
xs = [0.35, 2.20, 4.05, 5.90, 7.75]
for (txt, col), x in zip(STEPS, xs):
    box(s, x, 1.85, 1.65, 1.00, txt, fill=col, line=None, size=10, color=WHITE)
for gx in [2.02, 3.87, 5.72, 7.57]:
    arrow(s, gx, 2.20, 0.16, 0.30, color=DARK)
label(s, 0.35, 2.95, 9.05, 0.32,
      "→  live at  yourname-hello-staging.ml-capstone.cs.byu.edu", size=12, bold=True,
      color=ORANGE, align=PP_ALIGN.CENTER)

box(s, 0.35, 3.45, 9.05, 0.62,
    "If the tests fail at step 2, steps 3–5 never run.  Nothing deploys.  The old container keeps serving.",
    fill=None, line=RED, size=11.5, color=RED)

STEPS2 = [("6\nopen a PR\nstaging → main", BLUE), ("7\ntests run\nagain", PURPLE),
          ("8\nmerge", GREEN), ("9\nsame deploy,\nprod webhook", RED),
          ("10\nproduction\nswapped", GREEN)]
for (txt, col), x in zip(STEPS2, xs):
    box(s, x, 4.30, 1.65, 1.00, txt, fill=col, line=None, size=10, color=WHITE)
for gx in [2.02, 3.87, 5.72, 7.57]:
    arrow(s, gx, 4.65, 0.16, 0.30, color=DARK)
label(s, 0.35, 5.40, 9.05, 0.32,
      "→  live at  yourname-hello.ml-capstone.cs.byu.edu", size=12, bold=True,
      color=GREEN, align=PP_ALIGN.CENTER)

label(s, 0.35, 5.95, 9.05, 0.90,
      "You do steps 1, 6 and 8.  Everything else is automation you configured once.\n"
      "The review in step 6 is the last human checkpoint before users are affected.",
      size=12, italic=True, color=DARK, align=PP_ALIGN.CENTER)

bullets("The Pipeline Is Three Jobs", [
    "test — runs on every push and every PR",
    "deploy-staging — only on a push to staging, only if tests passed",
    "deploy-prod — only on a push to main, only if tests passed",
    "Written once in .github/workflows/ci.yml, in your repo",
    "The branch decides which environment. Nothing else does.",
])

code_slide("What That Looks Like in YAML", [
    "jobs:",
    "  test:",
    "    runs-on: ubuntu-latest",
    "    steps: [ checkout, install deps, pytest ]",
    "",
    "  deploy-staging:",
    "    needs: test                                  # <- the gate",
    "    if: github.ref == 'refs/heads/staging'       # <- the branch rule",
    "    steps:",
    "      - run: curl -X POST \"${{ secrets.COOLIFY_DEPLOY_WEBHOOK_STAGING }}\"",
    "",
    "  deploy-prod:",
    "    needs: test",
    "    if: github.ref == 'refs/heads/main'",
    "    steps:",
    "      - run: curl -X POST \"${{ secrets.COOLIFY_DEPLOY_WEBHOOK_PROD }}\"",
], size=13,
   caption="Two lines carry the whole policy: `needs: test` and the `if:` branch rule.")

bullets("Why a Webhook and a Secret", [
    "GitHub has to tell Coolify 'deploy now' — over the public internet",
    "The webhook URL is the instruction; the token proves it is you",
    "Both live in GitHub repository secrets, never in the repo",
    "A secret in git is a secret you have given away",
    ("rotate it, do not hide it — history is forever", 1),
])

bullets("The Deploy Is a Container Swap", [
    "Coolify builds a new image from your commit",
    "Starts new containers alongside the old ones",
    "Polls /health until it answers",
    "Only then does traffic move",
    "If /health never answers, the old container keeps serving",
    "That is your safety net, and you wrote it",
])

bullets("Health Checks Earn Their Keep", [
    "/health is not decoration — it is the deploy gate",
    "A container that starts is not the same as an app that works",
    "Make it check what would actually break:",
    ("can I reach the database?", 1),
    ("can I reach the service I depend on?", 1),
    "Return non-200 and the bad build never takes over",
])

bullets("Three Places Things Get Tested", [
    "Before you push — ./smoke-test.sh on your laptop, seconds",
    "In CI — unit tests on every push, no infrastructure needed",
    "After deploy — /health against the real environment",
    "Each catches what the previous one cannot",
    "None of them is optional if you want to sleep",
])

bullets("When It Breaks", [
    "Tests red → nothing deployed, fix and push again",
    "Deploy unhealthy → old container still serving, you have time",
    "Bad code already live → revert the commit and push",
    ("the pipeline that shipped it also un-ships it", 1),
    "Every deploy is traceable to a commit, and every commit to a person",
])

# ===========================================================================
# PART C — IaC
# ===========================================================================
section("Infrastructure as Code",
        "The part where you stop clicking")

bullets("How We Got Here", [
    "Somebody had to create the project, the environments,",
    "the two applications, the domains, the webhooks, the secrets",
    "In a UI. Roughly thirty clicks.",
    "Now do it for thirty students. Then do it again next semester.",
    "Then try to remember exactly what you clicked.",
])

bullets("Infrastructure as Code", [
    "Managing and provisioning infrastructure — servers, networks, databases —",
    "using machine-readable definition files rather than manual configuration",
    "Infrastructure becomes something you can review, version, and re-run",
    "Same tools you already use for code: git, pull requests, diffs",
])

bullets("Why It Matters", [
    "Reproducible — the same file gives the same result, every time",
    "Reviewable — a change to infrastructure shows up in a diff",
    "Recoverable — the environment is described, so it can be rebuilt",
    "Documented — the file IS the documentation, and it cannot go stale",
    "Scalable — thirty students is the same effort as one",
])

bullets("Imperative vs Declarative", [
    "Imperative: do this, then that, then this other thing",
    ("a script. Run it twice and you get two of everything.", 1),
    "Declarative: here is what should exist",
    ("the tool compares, then changes only what differs", 1),
    "Run it twice and nothing happens the second time",
    "That property has a name: idempotent",
])

bullets("Terraform", [
    "A general-purpose engine for declarative infrastructure",
    "You describe resources; it figures out how to get there",
    "Providers translate your description into API calls",
    ("AWS, GCP, Azure, GitHub, Cloudflare, Coolify, Docker…", 1),
    "Same workflow no matter what is underneath",
])

# --- terraform mental model diagram ---------------------------------------
s = prs.slides.add_slide(L_TONLY)
s.shapes.title.text = "Terraform's Four Ideas"

items = [
    ("PROVIDER", "the plugin that knows\nhow to talk to a platform", BLUE),
    ("RESOURCE", "one thing that should exist\n— a project, an app, a secret", GREEN),
    ("STATE", "what Terraform believes\nit already created", ORANGE),
    ("PLAN → APPLY", "compare desired vs actual,\nthen change only the difference", RED),
]
for i, (ttl, body, col) in enumerate(items):
    x = 0.40 + i * 2.35
    box(s, x, 1.80, 2.15, 0.52, ttl, fill=col, line=None, size=12, bold=True, color=WHITE)
    box(s, x, 2.38, 2.15, 1.15, body, fill=WHITE, line=col, size=10)

label(s, 0.40, 3.75, 9.05, 0.40,
      "Your .tf files describe the desired state.  State records the actual state.",
      size=13, bold=True, color=DARK, align=PP_ALIGN.CENTER)

box(s, 1.30, 4.30, 2.30, 0.62, "your .tf files\n(desired)", fill=GREEN, line=None,
    size=11, color=WHITE)
arrow(s, 3.75, 4.48, 0.55, 0.28, color=DARK)
box(s, 4.45, 4.30, 2.10, 0.62, "terraform plan", fill=RED, line=None, size=11, color=WHITE)
arrow(s, 6.70, 4.48, 0.55, 0.28, color=DARK)
box(s, 7.40, 4.30, 2.30, 0.62, "the platform\n(actual)", fill=ORANGE, line=None,
    size=11, color=WHITE)

label(s, 0.40, 5.15, 9.05, 0.95,
      "plan tells you what would change, before anything changes.\n"
      "Reading the plan is the habit that separates confident from lucky.",
      size=13, italic=True, color=DARK, align=PP_ALIGN.CENTER)

bullets("The Workflow", [
    "terraform init — download the providers named in your config",
    "terraform plan — show me what you would do",
    "terraform apply — do it",
    "terraform destroy — remove everything this config created",
    "You will run plan far more often than apply",
])

code_slide("Your main.tf — the Provider Block", [
    "terraform {",
    "  required_providers {",
    "    coolify = { source = \"bindtech-xyz/coolify\", version = \"~> 0.1.0\" }",
    "    github  = { source = \"integrations/github\",  version = \"~> 6.0\"  }",
    "  }",
    "}",
    "",
    "provider \"coolify\" {",
    "  endpoint = var.coolify_endpoint",
    "  token    = var.coolify_token      # never a literal token",
    "}",
    "",
    "provider \"github\" {",
    "  owner = var.github_org",
    "  token = var.github_token",
    "}",
], size=13,
   caption="Two providers, because this deploy touches two platforms. Versions are pinned so everyone gets the same behavior.")

code_slide("Variables — the Inputs", [
    "variable \"coolify_server_uuid\" {",
    "  description = \"UUID of YOUR team's Coolify server. Each team has its own.\"",
    "  type        = string",
    "",
    "  validation {",
    "    condition     = length(trimspace(var.coolify_server_uuid)) > 0",
    "    error_message = \"Required. Look it up: curl .../api/v1/servers\"",
    "  }",
    "}",
    "",
    "variable \"repo_name\" {",
    "  description = \"Just the repo name, not the org prefix\"",
    "  type        = string",
    "}",
], size=13,
   caption="Variables are the parameters of your infrastructure. Validation turns a confusing API error into a sentence that tells you what to do.")

code_slide("Resources — the Things That Should Exist", [
    "resource \"coolify_project\" \"app\" {",
    "  name = var.repo_name",
    "}",
    "",
    "resource \"coolify_environment\" \"staging\" {",
    "  project_uuid = coolify_project.app.uuid      # <- reference, not a copy",
    "  name         = \"staging\"",
    "}",
    "",
    "resource \"coolify_application\" \"staging\" {",
    "  environment_uuid       = coolify_environment.staging.uuid",
    "  git_repository         = \"${var.github_org}/${var.repo_name}\"",
    "  git_branch             = \"staging\"",
    "  build_pack             = \"dockercompose\"",
    "  is_auto_deploy_enabled = false                # GitHub Actions deploys",
    "}",
], size=12.5,
   caption="Each resource references the one it needs. Those references are what let Terraform work out the order by itself.")

# --- dependency graph ------------------------------------------------------
s = prs.slides.add_slide(L_TONLY)
s.shapes.title.text = "You Never Write the Order"

box(s, 3.85, 1.75, 2.30, 0.55, "coolify_project", fill=BLUE, line=None, size=11, color=WHITE)
arrow(s, 4.85, 2.35, 0.30, 0.30, color=DARK, shape=MSO_SHAPE.DOWN_ARROW)
box(s, 2.20, 2.72, 2.30, 0.55, "environment\nstaging", fill=ORANGE, line=None, size=10, color=WHITE)
box(s, 5.50, 2.72, 2.30, 0.55, "environment\nproduction", fill=GREEN, line=None, size=10, color=WHITE)
arrow(s, 3.20, 3.32, 0.30, 0.30, color=DARK, shape=MSO_SHAPE.DOWN_ARROW)
arrow(s, 6.50, 3.32, 0.30, 0.30, color=DARK, shape=MSO_SHAPE.DOWN_ARROW)
box(s, 2.20, 3.69, 2.30, 0.55, "application\nstaging", fill=ORANGE, line=None, size=10, color=WHITE)
box(s, 5.50, 3.69, 2.30, 0.55, "application\nproduction", fill=GREEN, line=None, size=10, color=WHITE)
arrow(s, 3.20, 4.29, 0.30, 0.30, color=DARK, shape=MSO_SHAPE.DOWN_ARROW)
arrow(s, 6.50, 4.29, 0.30, 0.30, color=DARK, shape=MSO_SHAPE.DOWN_ARROW)
box(s, 1.60, 4.66, 3.50, 0.55, "github secret: staging webhook", fill=PURPLE, line=None,
    size=10, color=WHITE)
box(s, 5.20, 4.66, 3.50, 0.55, "github secret: prod webhook", fill=PURPLE, line=None,
    size=10, color=WHITE)

label(s, 0.40, 5.45, 9.05, 1.10,
      "Terraform reads the references between resources and builds a graph.\n"
      "It creates the project first because everything else needs its UUID —\n"
      "not because you listed it first.",
      size=13, color=DARK, align=PP_ALIGN.CENTER)

code_slide("Outputs — What You Get Back", [
    "output \"staging_url\" {",
    "  value = \"http://${var.repo_name}-staging.${var.app_domain_base}\"",
    "}",
    "",
    "output \"production_app_uuid\" {",
    "  value = coolify_application.production.uuid",
    "}",
], size=13,
   caption="Outputs are the return values of your infrastructure — the facts you did not know until it existed.")

bullets("State: the Part People Trip Over", [
    "terraform.tfstate maps your resources to real IDs",
    "It is how Terraform knows 'I already made that'",
    "Delete it and Terraform forgets — then tries to create duplicates",
    "It can contain secrets. It is gitignored for a reason.",
    "Teams keep it in shared remote storage. You will keep it local.",
])

bullets("What It Cannot Do", [
    "Terraform is only as capable as its provider",
    "Our provider cannot set a per-service domain on a compose app",
    "So that one step stays manual — and is documented as manual",
    "Honest IaC: automate 90%, write down the 10%",
    "A fragile automation that fights the tool is worse than a checklist",
])

# --- portability -----------------------------------------------------------
s = prs.slides.add_slide(L_TONLY)
s.shapes.title.text = "Moving to a Different Provider"

box(s, 0.40, 1.75, 4.30, 0.50, "WHAT CHANGES", fill=RED, line=None, size=13, bold=True, color=WHITE)
box(s, 5.30, 1.75, 4.30, 0.50, "WHAT DOES NOT", fill=GREEN, line=None, size=13, bold=True, color=WHITE)

left = ["the provider block", "resource type names",
        "resource-specific arguments", "the credentials you supply"]
right = ["your application code", "your Dockerfile and compose file",
         "the CI workflow and its gates", "variables, outputs, structure",
         "init → plan → apply", "how you think about it"]
for i, t in enumerate(left):
    box(s, 0.40, 2.38 + i * 0.62, 4.30, 0.52, t, fill=WHITE, line=RED, size=11)
for i, t in enumerate(right):
    box(s, 5.30, 2.38 + i * 0.62, 4.30, 0.52, t, fill=WHITE, line=GREEN, size=11)

label(s, 0.40, 6.30, 9.20, 0.80,
      "You are not learning Coolify. You are learning a shape that AWS, GCP and Azure also have.",
      size=14, bold=True, color=DARK, align=PP_ALIGN.CENTER)

code_slide("The Same Idea, Somewhere Else", [
    "# today — Coolify on the class cluster",
    "resource \"coolify_application\" \"production\" {",
    "  environment_uuid = coolify_environment.production.uuid",
    "  git_repository   = \"${var.github_org}/${var.repo_name}\"",
    "  git_branch       = \"main\"",
    "}",
    "",
    "# tomorrow — AWS App Runner, same position in the same graph",
    "resource \"aws_apprunner_service\" \"production\" {",
    "  service_name = \"${var.repo_name}-prod\"",
    "  source_configuration {",
    "    image_repository { image_identifier = var.image_uri }",
    "  }",
    "}",
], size=12.5,
   caption="Different nouns, identical workflow. The GitHub resources and the whole CI pipeline do not move at all.")

bullets("Other Tools in This Space", [
    "AWS: CloudFormation, CDK",
    "Azure: ARM templates, Bicep",
    "GCP: Deployment Manager",
    "Kubernetes: manifests, Helm",
    "Pulumi: same idea, real programming languages",
    "OpenTofu: the open-source fork of Terraform — drop-in",
])

bullets("Docker Compose Is IaC Too", [
    "Declarative file describing services, networks, volumes",
    "compose up converges to the described state",
    "Same mental model, one machine instead of a cloud",
    "You have been writing IaC since the first lab",
    "Terraform does for the cloud what compose does for your laptop",
])

# ===========================================================================
# closing
# ===========================================================================
bullets("What You Will Actually Do", [
    "Use Claude / Codex / opencode to generate your .tf files",
    "That is fine — and it is why understanding the pieces matters",
    "You have to be able to read the plan and know if it is right",
    "You have to know which variable is wrong when it fails",
    "Generated code you cannot evaluate is a liability, not a shortcut",
])

bullets("The Questions to Ask Your Generated File", [
    "Which providers, and are the versions pinned?",
    "What resources does it create, and do I want all of them?",
    "What does each resource reference — is the graph right?",
    "Are any secrets hard-coded instead of variables?",
    "Does terraform plan show exactly what I expect, and nothing else?",
])

bullets("Recap", [
    "Local Docker → GitHub → CI → Coolify → staging → production → users",
    "Tests gate every deploy; health checks gate every swap",
    "Branch decides environment; nothing deploys by hand",
    "Infrastructure is described in a file, reviewed like code",
    "The provider changes. The shape does not.",
])

title_slide("Questions", "Then: build it")

prs.save(OUT)
print(f"saved {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
